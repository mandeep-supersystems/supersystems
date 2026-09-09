import sys, os
import uuid
from datetime import datetime
from decimal import Decimal
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, os.path.abspath("."))
from app import create_app, db

EXCEL_FILE = "inventory 08-09-2026.xlsx"
SHEET_NAME = "Production Bin"
TENANT_ID = "b424df0e-f766-4e94-b3fd-05777e158958"
WAREHOUSE_CODE = "MAIN"
REF_NO = "INV-08-09-2026"

app = create_app()

def run_import():
    print(f"Starting Inventory Import from '{EXCEL_FILE}'...")
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet '{SHEET_NAME}' not found in {EXCEL_FILE}")
    ws = wb[SHEET_NAME]

    # Read items
    items = []
    for r in range(3, ws.max_row + 1):
        pn = ws.cell(r, 1).value
        sspn = ws.cell(r, 2).value
        mpn = ws.cell(r, 3).value
        make = ws.cell(r, 4).value
        qty = ws.cell(r, 5).value
        bin_val = ws.cell(r, 6).value

        if not any([pn, sspn, mpn, make, qty, bin_val]):
            continue

        items.append({
            "row": r,
            "part_number": str(pn).strip() if pn else None,
            "sspn": str(sspn).strip() if sspn else "",
            "mpn": str(mpn).strip() if mpn else None,
            "make": str(make).strip() if make else "",
            "qty": Decimal(str(qty)) if qty is not None else Decimal(0),
            "bin": str(bin_val).strip() if bin_val else None
        })

    print(f"Read {len(items)} items from workbook.")

    with app.app_context():
        # Step 1: Ensure mpn column exists in inventory_stock_levels and inventory_stock_movements
        print("\n--- Step 1: Checking schema for mpn columns ---")
        db.session.execute(db.text("ALTER TABLE inventory_stock_levels ADD COLUMN IF NOT EXISTS mpn VARCHAR(255);"))
        db.session.execute(db.text("ALTER TABLE inventory_stock_movements ADD COLUMN IF NOT EXISTS mpn VARCHAR(255);"))
        db.session.commit()
        print("Schema ensured (mpn columns present).")

        # Step 2: Ensure bins E80, E81, E82 exist in inventory_locations and public.warehouse_bins
        print("\n--- Step 2: Checking/creating bin locations ---")
        bins_to_create = {
            "E80": {"row": "Row 8", "col": "Col 0", "code": "P134-2F-E-E80"},
            "E81": {"row": "Row 8", "col": "Col 1", "code": "P134-2F-E-E81"},
            "E82": {"row": "Row 8", "col": "Col 2", "code": "P134-2F-E-E82"},
        }

        # Calculate total units per bin from items
        bin_unit_totals = {}
        for it in items:
            b = it["bin"]
            bin_unit_totals[b] = bin_unit_totals.get(b, Decimal(0)) + it["qty"]

        for bcode, bmeta in bins_to_create.items():
            # Check inventory_locations
            loc_exists = db.session.execute(
                db.text("SELECT id FROM inventory_locations WHERE bin_code = :b AND (tenant_id = :tid OR tenant_id = 'TEST')"),
                {"b": bcode, "tid": TENANT_ID}
            ).first()

            current_occ = bin_unit_totals.get(bcode, Decimal(0))

            if not loc_exists:
                loc_id = str(uuid.uuid4())
                db.session.execute(db.text("""
                    INSERT INTO inventory_locations (
                        id, location_code, plant, floor_name, shelf_name, row_name, column_name,
                        bin_code, warehouse_code, capacity, current_occupancy, is_active, is_deleted,
                        tenant_id, created_at
                    ) VALUES (
                        :id, :loc_code, '134', '2nd Floor', 'Shelf E', :row, :col,
                        :bcode, :wh, 500, :occ, true, false,
                        :tid, CURRENT_TIMESTAMP
                    )
                """), {
                    "id": loc_id, "loc_code": bmeta["code"], "row": bmeta["row"], "col": bmeta["col"],
                    "bcode": bcode, "wh": WAREHOUSE_CODE, "occ": float(current_occ), "tid": TENANT_ID
                })
                print(f"  Created inventory_locations entry for bin {bcode} ({bmeta['code']})")
            else:
                db.session.execute(db.text("""
                    UPDATE inventory_locations SET current_occupancy = current_occupancy + :occ
                    WHERE bin_code = :bcode AND (tenant_id = :tid OR tenant_id = 'TEST')
                """), {"occ": float(current_occ), "bcode": bcode, "tid": TENANT_ID})
                print(f"  Updated inventory_locations occupancy for bin {bcode}")

            # Check public.warehouse_bins
            wh_bin_exists = db.session.execute(
                db.text("SELECT id FROM public.warehouse_bins WHERE bin_code = :b AND tenant_id = :tid"),
                {"b": bcode, "tid": TENANT_ID}
            ).first()

            if not wh_bin_exists:
                wh_bin_id = str(uuid.uuid4())
                db.session.execute(db.text("""
                    INSERT INTO public.warehouse_bins (
                        id, bin_code, zone_code, warehouse_code, aisle, rack, level,
                        capacity_units, current_units, status, is_deleted, tenant_id,
                        created_at, updated_at, location_code, bin_type, bin_color
                    ) VALUES (
                        :id, :bcode, 'Shelf E', :wh, 'Shelf E', :rack, :level,
                        500, :units, 'active', false, :tid,
                        CURRENT_TIMESTAMP, CURRENT_TIMESTAMP, :loc_code, 'medium', ''
                    )
                """), {
                    "id": wh_bin_id, "bcode": bcode, "wh": WAREHOUSE_CODE, "rack": bmeta["row"], "level": bmeta["col"],
                    "units": int(current_occ), "tid": TENANT_ID, "loc_code": bmeta["code"]
                })
                print(f"  Created public.warehouse_bins entry for bin {bcode}")
            else:
                db.session.execute(db.text("""
                    UPDATE public.warehouse_bins SET current_units = current_units + :units, updated_at = CURRENT_TIMESTAMP
                    WHERE bin_code = :bcode AND tenant_id = :tid
                """), {"units": int(current_occ), "bcode": bcode, "tid": TENANT_ID})
                print(f"  Updated public.warehouse_bins current_units for bin {bcode}")

        db.session.commit()

        # Step 3: Register missing MPNs in part.manufacturers
        print("\n--- Step 3: Registering missing MPNs in part.manufacturers ---")
        existing_mfrs = db.session.execute(db.text("SELECT part_number, mpn FROM part.manufacturers")).fetchall()
        existing_mfr_set = {(r[0].strip().upper() if r[0] else "", r[1].strip().upper() if r[1] else "") for r in existing_mfrs}

        new_mfr_count = 0
        for it in items:
            pn_key = it["part_number"].upper() if it["part_number"] else ""
            mpn_key = it["mpn"].upper() if it["mpn"] else ""
            if (pn_key, mpn_key) not in existing_mfr_set and mpn_key:
                mfr_id = str(uuid.uuid4())
                db.session.execute(db.text("""
                    INSERT INTO part.manufacturers (id, part_number, mpn, make, created_at, updated_at)
                    VALUES (:id, :pn, :mpn, :make, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
                """), {
                    "id": mfr_id,
                    "pn": it["part_number"],
                    "mpn": it["mpn"],
                    "make": it["make"]
                })
                existing_mfr_set.add((pn_key, mpn_key))
                new_mfr_count += 1

        db.session.commit()
        print(f"Registered {new_mfr_count} new manufacturer/MPN mappings in part.manufacturers.")

        # Step 4: Insert items into inventory_stock_levels and record stock movements
        print("\n--- Step 4: Inserting into inventory_stock_levels & inventory_stock_movements ---")
        inserted_stock = 0
        inserted_movements = 0

        for it in items:
            stock_id = str(uuid.uuid4())
            bmeta = bins_to_create.get(it["bin"], {"code": f"P134-2F-E-{it['bin']}"})
            loc_code = bmeta["code"]
            desc = it["sspn"] if it["sspn"] else ""

            # Insert stock level
            db.session.execute(db.text("""
                INSERT INTO inventory_stock_levels (
                    id, part_number, part_description, item_type, warehouse_id, warehouse_code,
                    zone_code, bin_code, qty_on_hand, qty_reserved, qty_available,
                    reorder_point, reorder_qty, unit, unit_cost, total_value,
                    last_movement_at, is_deleted, tenant_id, created_at, updated_at,
                    location_code, manufacturer, mpn
                ) VALUES (
                    :id, :pn, :desc, 'PART', NULL, :wh,
                    'Shelf E', :bin, :qty, 0, :qty,
                    0, 0, 'pcs', 0, 0,
                    CURRENT_TIMESTAMP, false, :tid, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP,
                    :loc_code, :make, :mpn
                )
            """), {
                "id": stock_id,
                "pn": it["part_number"],
                "desc": desc,
                "wh": WAREHOUSE_CODE,
                "bin": it["bin"],
                "qty": it["qty"],
                "tid": TENANT_ID,
                "loc_code": loc_code,
                "make": it["make"],
                "mpn": it["mpn"]
            })
            inserted_stock += 1

            # Insert stock movement audit log
            mov_id = str(uuid.uuid4())
            mov_no = f"MOV-{datetime.now().strftime('%Y%m%d%H%M%S')}-{it['row']}"
            db.session.execute(db.text("""
                INSERT INTO inventory_stock_movements (
                    id, movement_no, movement_type, part_number, part_description, item_type,
                    from_warehouse_code, from_bin_code, to_warehouse_code, to_bin_code,
                    qty, unit, unit_cost, reference_type, reference_no, reason,
                    performed_by, tenant_id, created_at, mpn
                ) VALUES (
                    :id, :mov_no, 'receipt', :pn, :desc, 'PART',
                    '-', '-', :to_wh, :to_bin,
                    :qty, 'pcs', 0, 'PHYSICAL_INVENTORY', :ref_no, 'Physical Inventory Import 08-09-2026',
                    'System (Excel Import)', :tid, CURRENT_TIMESTAMP, :mpn
                )
            """), {
                "id": mov_id,
                "mov_no": mov_no,
                "pn": it["part_number"],
                "desc": desc,
                "to_wh": WAREHOUSE_CODE,
                "to_bin": it["bin"],
                "qty": it["qty"],
                "ref_no": REF_NO,
                "tid": TENANT_ID,
                "mpn": it["mpn"]
            })
            inserted_movements += 1

        db.session.commit()
        print(f"Successfully inserted {inserted_stock} stock level records into inventory_stock_levels.")
        print(f"Successfully recorded {inserted_movements} movement records into inventory_stock_movements.")
        print("\n=== IMPORT COMPLETE ===")

if __name__ == "__main__":
    run_import()
